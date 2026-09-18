"""
Document processor — reads and summarizes uploaded support documents.
Handles PDF, Office (.docx/.xlsx), CSV and any text-like format; binary files that
have no parser are reported rather than decoded into noise.
Integrates uploaded documents into BRD analysis.
"""

import json
from pathlib import Path
from typing import Dict, List, Any, Optional

# Extensions read straight through as UTF-8 text.
_TEXT_SUFFIXES = {
    ".md", ".markdown", ".txt", ".text", ".json", ".yaml", ".yml", ".xml",
    ".html", ".htm", ".rst", ".log", ".ini", ".cfg", ".toml", ".sql",
    ".java", ".py", ".js", ".ts", ".properties", ".gherkin", ".feature",
}


class DocumentProcessor:
    """Process uploaded supporting documents for BRD analysis."""

    def __init__(self, kb_path: Path):
        """Initialize with KB directory path."""
        self.kb_path = Path(kb_path)
        self.docs_dir = self.kb_path / "_uploaded_documents"

    def has_documents(self) -> bool:
        """Check if uploaded documents directory exists and has files."""
        if not self.docs_dir.is_dir():
            return False
        return any(self.docs_dir.iterdir())

    def get_all_documents(self) -> List[Path]:
        """Get all uploaded document files."""
        if not self.docs_dir.is_dir():
            return []
        return sorted([f for f in self.docs_dir.iterdir() if f.is_file()])

    def read_document(self, doc_path: Path) -> Optional[str]:
        """Read and return document content."""
        suffix = doc_path.suffix.lower()
        try:
            if suffix == ".pdf":
                return self._read_pdf(doc_path)
            if suffix == ".csv":
                return self._read_csv(doc_path)
            if suffix == ".docx":
                return self._read_docx(doc_path)
            if suffix in {".xlsx", ".xlsm"}:
                return self._read_xlsx(doc_path)
            if suffix in _TEXT_SUFFIXES:
                return self._read_text(doc_path)
            return self._read_unknown(doc_path)
        except Exception as e:
            print(f"Warning: Could not read document {doc_path.name}: {e}")
            return None

    def _read_text(self, path: Path) -> str:
        """Read plain text file."""
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            return ""

    def _read_unknown(self, path: Path) -> str:
        """Read an unrecognised extension, refusing anything that looks binary.

        Decoding a binary file with errors="replace" used to push pages of
        mojibake into the prompt, which costs tokens and tells the model nothing.
        """
        try:
            head = path.read_bytes()[:8192]
        except Exception:
            return ""

        if b"\x00" in head:
            return f"[Binary file: {path.name} — no parser available, content not included]"

        return self._read_text(path)

    def _read_csv(self, path: Path) -> str:
        """Read CSV file and format as table."""
        try:
            import csv

            content = path.read_text(encoding="utf-8", errors="replace")
            lines = content.split("\n")

            # Parse CSV
            reader = csv.reader(lines)
            rows = list(reader)

            if not rows:
                return ""

            # Format as markdown table
            header = rows[0]
            markdown = "| " + " | ".join(header) + " |\n"
            markdown += "|" + "|".join(["---"] * len(header)) + "|\n"

            for row in rows[1:]:
                if row:
                    markdown += "| " + " | ".join(row) + " |\n"

            return markdown
        except Exception:
            return ""

    def _read_pdf(self, path: Path) -> str:
        """Read PDF file (text extraction)."""
        try:
            import pdfplumber

            text = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text.append(page.extract_text())
            return "\n".join(text)
        except ImportError:
            print("Warning: pdfplumber not installed. Cannot read PDF files.")
            return f"[PDF Document: {path.name}]\n(pdfplumber not installed - please install to extract PDF text)"
        except Exception as e:
            return f"[PDF Document: {path.name}]\n(Error reading PDF: {e})"

    def _read_docx(self, path: Path) -> str:
        """Read a Word document: paragraphs plus table cells."""
        try:
            import docx
        except ImportError:
            return f"[Word Document: {path.name}]\n(python-docx not installed - cannot extract text)"

        try:
            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append("| " + " | ".join(cells) + " |")
            return "\n".join(parts)
        except Exception as e:
            return f"[Word Document: {path.name}]\n(Error reading document: {e})"

    def _read_xlsx(self, path: Path) -> str:
        """Read a spreadsheet, one markdown table per sheet."""
        try:
            import openpyxl
        except ImportError:
            return f"[Spreadsheet: {path.name}]\n(openpyxl not installed - cannot extract cells)"

        try:
            book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for sheet in book.worksheets:
                rows = [
                    ["" if c is None else str(c) for c in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [r for r in rows if any(v.strip() for v in r)]
                if not rows:
                    continue
                parts.append(f"\n### Sheet: {sheet.title}\n")
                parts.append("| " + " | ".join(rows[0]) + " |")
                parts.append("|" + "|".join(["---"] * len(rows[0])) + "|")
                for row in rows[1:]:
                    parts.append("| " + " | ".join(row) + " |")
            book.close()
            return "\n".join(parts)
        except Exception as e:
            return f"[Spreadsheet: {path.name}]\n(Error reading spreadsheet: {e})"

    def get_documents_context(
        self,
        max_docs: int = 25,
        max_total_len: int = 60000,
        max_doc_len: int = 20000,
    ) -> str:
        """Get formatted context string for all uploaded documents.

        Args:
            max_docs: Maximum number of documents to include
            max_total_len: Maximum total character length to prevent token overflow
            max_doc_len: Maximum characters taken from any single document
        """
        if not self.has_documents():
            return ""

        docs = self.get_all_documents()[:max_docs]
        context_parts = ["# 📎 Uploaded Supporting Documents (User-Provided Requirements)\n"]
        context_parts.append("⭐ These documents contain critical business requirements and specifications\n")

        total_len = 0
        included_docs = 0

        for idx, doc in enumerate(docs, 1):
            content = self.read_document(doc)
            if content:
                # Calculate space needed for this document
                doc_header = f"\nDocument {idx}: {doc.name}\n{'='*60}\n"
                doc_footer = f"\n{'='*60}\n"
                space_needed = len(doc_header) + len(doc_footer)

                if len(content) > max_doc_len:
                    content = content[:max_doc_len] + "\n[... truncated ...]"

                # Check if adding this document would exceed total limit
                if total_len + space_needed + len(content) > max_total_len:
                    context_parts.append(f"\n[... {len(docs) - included_docs} additional document(s) not included due to size limit ...]")
                    break

                context_parts.append(doc_header)
                context_parts.append(content)
                context_parts.append(doc_footer)

                total_len += space_needed + len(content)
                included_docs += 1

        if included_docs > 0:
            context_parts.append(f"\n✓ Included {included_docs} document(s) for analysis")

        return "\n".join(context_parts)

    def get_documents_summary(self) -> Dict[str, Any]:
        """Get summary of uploaded documents."""
        if not self.has_documents():
            return {
                "has_documents": False,
                "document_count": 0,
                "documents": []
            }

        docs = self.get_all_documents()
        doc_info = []

        for doc in docs:
            try:
                size = doc.stat().st_size
                doc_info.append({
                    "name": doc.name,
                    "type": doc.suffix.lower(),
                    "size": size,
                    "size_kb": round(size / 1024, 2)
                })
            except Exception:
                pass

        return {
            "has_documents": True,
            "document_count": len(doc_info),
            "documents": doc_info
        }
